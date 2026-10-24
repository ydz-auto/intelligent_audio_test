# -*- coding: utf-8 -*-
"""INT-1624 验收（测试工程师独立脚本）——接口契约层。

对 docker 真实库栈网关 http://127.0.0.1:15000（库 intelligent_audio_test）直接核对
INT-816 三资源 export/import 端点契约，作为前端实测的对照基线：

  * 导出 JSON / Excel：字段白名单逐字段、不含 id / 时间戳字段、Content-Disposition 命名；
  * 导出 ids 过滤（= 弹窗「当前页选中 / 当前筛选结果」范围的底层契约）；
  * 导入闭环：新建 → 再导入命中匹配键走更新，不产生第二条；
  * 部分成功：合法行 + 缺匹配键行 → updated=1 / skipped=1，errors 带「第N行」；
  * 异常：.txt / 空文件 / 非法 JSON → 400 且有中文 message；
  * 三资源（test-devices / playback-devices / apis）各走一次。

数据纪律：只写「接线测试_INT1624_」前缀记录，脚本末尾删除并复核零残留。
"""
import functools
import io
import json
import sys
import time
import zipfile

import requests

print = functools.partial(print, flush=True)  # noqa: A001
try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

BASE = 'http://127.0.0.1:15000/api/v1'
PREFIX = '接线测试_INT1624_'
SERIAL_PREFIX = 'JX1624SN'

TEST_DEVICE_FIELDS = [
    'name', 'model', 'type', 'system', 'system_version', 'app_name',
    'app_version', 'description', 'location', 'max_audio_duration',
    'needs_prompt_audio', 'prompt_config', 'connection_type',
    'keywords', 'serial_number', 'ip', 'status', 'supported_algorithms',
]
PLAYBACK_FIELDS = [
    'name', 'model', 'device_type', 'sample_rate', 'device_unique_id',
    'channel_index', 'description', 'status',
]
API_FIELDS = [
    'name', 'vendor', 'description', 'meta', 'algorithm_type',
    'default_max_process', 'default_max_timeout', 'default_max_audio_duration',
    'status', 'endpoints', 'output_types', 'device_type', 'adapter_class',
    'audio_config',
]

FAILS = []
CREATED = {'test': [], 'playback': [], 'api': []}


def check(name, cond, detail=''):
    print(('[PASS] ' if cond else '[FAIL] ') + name + ((' — ' + str(detail)) if detail else ''))
    if not cond:
        FAILS.append(name)


def get_json(path):
    r = requests.get(BASE + path, timeout=120)
    return r.status_code, r.json() if r.headers.get('content-type', '').startswith('application/json') else None


def items(path):
    status, payload = get_json(path)
    return ((payload or {}).get('data') or {}).get('items') or []


def post_import(resource, filename, content, mime='application/json'):
    files = {'file': (filename, content, mime)}
    r = requests.post(BASE + '/' + resource + '/import', files=files, timeout=300)
    try:
        return r.status_code, r.json()
    except Exception:
        return r.status_code, None


def export(resource, fmt='json', ids=None):
    params = {'format': fmt}
    if ids:
        params['ids'] = ','.join(str(i) for i in ids)
    r = requests.get(BASE + '/' + resource + '/export', params=params, timeout=180)
    return r


def cleanup():
    """删除本卡自建数据并复核零残留。"""
    removed = {'test': 0, 'playback': 0, 'api': 0}
    for it in items('/test-devices?page=1&per_page=1000'):
        name = str(it.get('name') or '')
        serial = str(it.get('serial_number') or '')
        if name.startswith(PREFIX) or serial.startswith(SERIAL_PREFIX):
            requests.delete(BASE + '/test-devices/%s' % it['id'], timeout=180)
            removed['test'] += 1
    for it in items('/playback-devices?page=1&per_page=1000'):
        if str(it.get('name') or '').startswith(PREFIX) or str(it.get('device_unique_id') or '').startswith(SERIAL_PREFIX):
            requests.delete(BASE + '/playback-devices/%s' % it['id'], timeout=180)
            removed['playback'] += 1
    for it in items('/apis?page=1&per_page=1000'):
        if str(it.get('name') or '').startswith(PREFIX):
            requests.delete(BASE + '/apis/%s' % it['id'], timeout=180)
            removed['api'] += 1
    time.sleep(2)
    left = []
    for res in ('test-devices', 'playback-devices', 'apis'):
        for it in items('/' + res + '?page=1&per_page=1000'):
            if str(it.get('name') or '').startswith(PREFIX):
                left.append('%s:%s' % (res, it['id']))
    return removed, left


def main():
    t0 = time.time()
    st, _ = get_json('/test-devices?page=1&per_page=1')
    print('gateway probe status =', st)

    # 预留：清掉上一轮残留
    removed, left = cleanup()
    print('pre-clean removed', removed, 'left', left)

    # ---------- 1. 导出 JSON：白名单 ----------
    r = export('test-devices')
    check('导出 JSON 200', r.status_code == 200, r.status_code)
    cd = r.headers.get('content-disposition', '')
    check('Content-Disposition 文件名 test_devices_<ts>.json',
          cd.startswith('attachment;') and 'test_devices_' in cd and cd.strip().endswith('.json'), cd)
    payload = json.loads(r.content.decode('utf-8'))
    rows = payload.get('test_devices') or []
    check('JSON 信封含 test_devices / exported_at / total_count',
          isinstance(rows, list) and 'exported_at' in payload and 'total_count' in payload,
          list(payload.keys()))
    check('total_count == 行数', payload.get('total_count') == len(rows), (payload.get('total_count'), len(rows)))
    if rows:
        extra = set(rows[0].keys()) - set(TEST_DEVICE_FIELDS)
        missing = set(TEST_DEVICE_FIELDS) - set(rows[0].keys())
        check('测试设备导出字段 == 白名单（无 id / 无时间戳 / 无告警字段）', not extra and not missing,
              'extra=%s missing=%s' % (sorted(extra), sorted(missing)))

    # ---------- 2. 导出 Excel：表头 = 白名单 ----------
    r = export('test-devices', 'excel')
    check('导出 Excel 200', r.status_code == 200, r.status_code)
    check('Excel 文件头 PK', r.content[:2] == b'PK', r.content[:8])
    cd = r.headers.get('content-disposition', '')
    check('Content-Disposition 扩展名 .xlsx', cd.strip().endswith('.xlsx'), cd)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        ws = wb[wb.sheetnames[0]]
        header = [c.value for c in next(ws.iter_rows(max_row=1))]
        check('Excel 表头 == 白名单', header == TEST_DEVICE_FIELDS, header)
    except Exception as exc:  # noqa: BLE001
        check('Excel 表头 == 白名单', False, 'openpyxl: %s' % exc)

    # ---------- 3. ids 过滤（选中 / 筛选范围契约） ----------
    base_rows = items('/test-devices?page=1&per_page=1000')
    if base_rows:
        target = base_rows[0]
        r = export('test-devices', 'json', ids=[target['id']])
        sub = (json.loads(r.content.decode('utf-8')).get('test_devices') or [])
        check('ids 过滤只返回指定 1 条', len(sub) == 1 and sub[0].get('name') == target.get('name'),
              'ids=%s got=%s' % ([target['id']], [x.get('name') for x in sub]))
        check('ids 过滤行字段与列表接口一致',
              sub and sub[0].get('serial_number') == target.get('serial_number'))

    # ---------- 4. 导入：新建 → 再导入更新（测试设备） ----------
    name = PREFIX + '设备A'
    serial = SERIAL_PREFIX + '001'
    doc = {'test_devices': [{
        'name': name, 'model': 'JX-Model', 'type': 'physical', 'system': 'Android',
        'system_version': '14', 'app_name': 'iatp', 'app_version': '1.0',
        'description': 'v1', 'location': 'JX-Lab', 'max_audio_duration': 30,
        'needs_prompt_audio': False, 'connection_type': 'usb',
        'keywords': 'jx,kw', 'serial_number': serial, 'ip': '10.0.0.9',
        'status': 'offline', 'supported_algorithms': [],
    }]}
    body = json.dumps(doc, ensure_ascii=False).encode('utf-8')
    st, res = post_import('test-devices', 'create.json', body)
    data = (res or {}).get('data') or {}
    check('导入新建 200 且 imported=1', st == 200 and data.get('imported_count') == 1, (st, res))
    check('导入 message 文案「成功导入 1 个，更新 0 个，跳过 0 个」',
          (res or {}).get('message') == '成功导入 1 个，更新 0 个，跳过 0 个', (res or {}).get('message'))
    rows_now = [it for it in items('/test-devices?page=1&per_page=1000') if it.get('name') == name]
    check('新建记录回读 1 条且字段一致',
          len(rows_now) == 1 and rows_now[0].get('serial_number') == serial
          and rows_now[0].get('location') == 'JX-Lab' and rows_now[0].get('model') == 'JX-Model',
          rows_now[:1])
    CREATED['test'].append(name)

    doc['test_devices'][0]['description'] = 'v2-updated'
    st, res = post_import('test-devices', 'update.json',
                          json.dumps(doc, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    check('再导入命中匹配键走更新 updated=1 / imported=0',
          st == 200 and data.get('updated_count') == 1 and data.get('imported_count') == 0, (st, data))
    rows_now = [it for it in items('/test-devices?page=1&per_page=1000') if it.get('name') == name]
    check('更新后仍只 1 条、description 已变（无重复脏数据）',
          len(rows_now) == 1 and rows_now[0].get('description') == 'v2-updated', rows_now[:1])

    # ---------- 5. 部分成功：合法行 + 缺匹配键行 ----------
    mixed = {'test_devices': [
        {'name': name, 'model': 'JX-Model', 'type': 'physical', 'system': 'Android',
         'system_version': '14', 'app_name': 'iatp', 'app_version': '1.0',
         'description': 'v3-mixed', 'serial_number': serial, 'status': 'offline'},
        {'name': PREFIX + '缺键行', 'model': 'JX-Model', 'type': 'physical',
         'system': 'Android', 'system_version': '14', 'app_name': 'iatp',
         'app_version': '1.0', 'description': 'missing-key'},
    ]}
    st, res = post_import('test-devices', 'mixed.json',
                          json.dumps(mixed, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    errs = data.get('errors') or []
    check('混合文件 200 且 updated=1 / skipped=1',
          st == 200 and data.get('updated_count') == 1 and data.get('skipped_count') == 1, (st, data))
    check('errors 含行号「第2行」', any('第2行' in str(e) for e in errs), errs)
    check('缺键行未落库（无重复脏数据）',
          not [it for it in items('/test-devices?page=1&per_page=1000')
               if it.get('name') == PREFIX + '缺键行'])

    # ---------- 6. 异常文件 ----------
    st, res = post_import('test-devices', 'bad.txt', b'not a supported file', 'text/plain')
    check('.txt → 400 且有中文提示', st == 400 and bool((res or {}).get('message')), (st, (res or {}).get('message')))
    st, res = post_import('test-devices', 'empty.json', b'')
    check('空文件 → 400 且有中文提示', st == 400 and bool((res or {}).get('message')), (st, (res or {}).get('message')))
    st, res = post_import('test-devices', 'invalid.json', b'{"test_devices": [oops')
    check('非法 JSON → 400 且有中文提示', st == 400 and bool((res or {}).get('message')), (st, (res or {}).get('message')))
    r = requests.post(BASE + '/test-devices/import', data={}, timeout=120)
    check('无文件 → 400', r.status_code == 400, r.status_code)

    # ---------- 7. 播放设备 ----------
    pb_name = PREFIX + '播放A'
    pb_uid = SERIAL_PREFIX + 'UID001'
    pb_doc = {'playback_devices': [{
        'name': pb_name, 'model': 'JX-PB', 'device_type': 'dry', 'sample_rate': 48000,
        'device_unique_id': pb_uid, 'channel_index': 0, 'description': 'pb-v1', 'status': 'offline',
    }]}
    st, res = post_import('playback-devices', 'pb.json',
                          json.dumps(pb_doc, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    check('播放设备导入新建 imported=1', st == 200 and data.get('imported_count') == 1, (st, data))
    pb_doc['playback_devices'][0]['description'] = 'pb-v2'
    st, res = post_import('playback-devices', 'pb2.json',
                          json.dumps(pb_doc, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    check('播放设备再导入 updated=1', st == 200 and data.get('updated_count') == 1, (st, data))
    pb_rows = [it for it in items('/playback-devices?page=1&per_page=1000') if it.get('name') == pb_name]
    check('播放设备仍 1 条（无重复）', len(pb_rows) == 1, len(pb_rows))
    CREATED['playback'].append(pb_name)
    r = export('playback-devices')
    pb_payload = json.loads(r.content.decode('utf-8'))
    pb_export = pb_payload.get('playback_devices') or []
    check('播放设备导出信封 + 字段白名单',
          r.status_code == 200 and pb_export
          and not (set(pb_export[0].keys()) - set(PLAYBACK_FIELDS))
          and not (set(PLAYBACK_FIELDS) - set(pb_export[0].keys())),
          sorted(pb_export[0].keys()) if pb_export else [])

    # ---------- 8. 测试 API ----------
    api_name = PREFIX + 'API-A'
    api_doc = {'apis': [{
        'name': api_name, 'vendor': 'jx-vendor', 'description': 'api-v1', 'status': 'online',
        'meta': {}, 'default_max_process': 5, 'default_max_timeout': 30,
        'default_max_audio_duration': 60, 'endpoints': [], 'output_types': [],
    }]}
    st, res = post_import('apis', 'api.json',
                          json.dumps(api_doc, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    check('测试 API 导入新建 imported=1', st == 200 and data.get('imported_count') == 1, (st, data))
    api_doc['apis'][0]['description'] = 'api-v2'
    st, res = post_import('apis', 'api2.json',
                          json.dumps(api_doc, ensure_ascii=False).encode('utf-8'))
    data = (res or {}).get('data') or {}
    check('测试 API 再导入 updated=1', st == 200 and data.get('updated_count') == 1, (st, data))
    api_rows = [it for it in items('/apis?page=1&per_page=1000') if it.get('name') == api_name]
    check('测试 API 仍 1 条（无重复）', len(api_rows) == 1, len(api_rows))
    CREATED['api'].append(api_name)
    r = export('apis')
    api_payload = json.loads(r.content.decode('utf-8'))
    api_export = api_payload.get('apis') or []
    check('测试 API 导出信封 + 字段白名单',
          r.status_code == 200 and api_export
          and not (set(api_export[0].keys()) - set(API_FIELDS))
          and not (set(API_FIELDS) - set(api_export[0].keys())),
          sorted(api_export[0].keys()) if api_export else [])

    # ---------- 7. 导出→删除→导入其导出文件（三资源各一次，验收标准 3） ----------
    def roundtrip(resource, list_path, envelope_key, doc, del_path_tmpl, name_key='name'):
        body = json.dumps(doc, ensure_ascii=False).encode('utf-8')
        st, res = post_import(resource, 'rt-create.json', body)
        ok_create = st == 200 and ((res or {}).get('data') or {}).get('imported_count') == 1
        rows = [it for it in items(list_path) if it.get('name') == doc[envelope_key][0][name_key]]
        check('%s 往返前置：导入新建 1 条' % resource, ok_create and len(rows) == 1, (st, (res or {}).get('message')))
        if not rows:
            return
        rid = rows[0]['id']
        # 导出这一条（ids 过滤，等价弹窗「选中 / 筛选」范围）
        r = export(resource, 'json', ids=[rid])
        doc_out = json.loads(r.content.decode('utf-8'))
        exported = doc_out.get(envelope_key) or []
        check('%s 导出该条 1 行' % resource, r.status_code == 200 and len(exported) == 1, len(exported))
        # 字段值逐个与提交值一致（白名单逐字段）
        src = doc[envelope_key][0]
        diffs = []
        for k, v in src.items():
            got = exported[0].get(k)
            if isinstance(v, (list, dict)) or v is None:
                continue
            if isinstance(v, float) and isinstance(got, (int, float)) and float(got) == float(v):
                continue
            if got != v:
                diffs.append('%s: %r != %r' % (k, got, v))
        check('%s 导出值逐字段与提交值一致' % resource, not diffs, diffs)
        # 删除
        dr = requests.delete(BASE + del_path_tmpl % rid, timeout=180)
        gone = not [it for it in items(list_path) if it.get('id') == rid]
        check('%s 删除后列表消失' % resource, dr.status_code < 300 and gone, dr.status_code)
        # 用导出的文件本体（不改动）导回
        st, res = post_import(resource, 'rt-roundtrip.json',
                              json.dumps(doc_out, ensure_ascii=False).encode('utf-8'))
        data = (res or {}).get('data') or {}
        back = [it for it in items(list_path) if it.get('name') == src[name_key]]
        check('%s 用导出文件导回：新建 1 条、GET 回读一致' % resource,
              st == 200 and data.get('imported_count') == 1 and len(back) == 1
              and back[0].get('model') == src.get('model'),
              (st, data, [(b.get('id'), b.get('name')) for b in back]))

    roundtrip('test-devices', '/test-devices?page=1&per_page=1000', 'test_devices', {
        'test_devices': [{'name': PREFIX + '往返', 'model': 'RT-Model', 'type': 'physical',
                          'system': 'Android', 'system_version': '14', 'app_name': 'iatp',
                          'app_version': '1.0', 'description': 'roundtrip', 'location': 'RT-Lab',
                          'max_audio_duration': 45, 'needs_prompt_audio': False,
                          'connection_type': 'usb', 'keywords': 'rt,kw',
                          'serial_number': SERIAL_PREFIX + 'RT1', 'ip': '10.0.0.11',
                          'status': 'offline', 'supported_algorithms': []}]},
        '/test-devices/%s')
    roundtrip('playback-devices', '/playback-devices?page=1&per_page=1000', 'playback_devices', {
        'playback_devices': [{'name': PREFIX + '播放往返', 'model': 'RT-PB', 'device_type': 'dry',
                              'sample_rate': 44100, 'device_unique_id': SERIAL_PREFIX + 'RTUID',
                              'channel_index': 1, 'description': 'rt-pb', 'status': 'offline'}]},
        '/playback-devices/%s')
    roundtrip('apis', '/apis?page=1&per_page=1000', 'apis', {
        'apis': [{'name': PREFIX + 'API往返', 'vendor': 'rt-vendor', 'description': 'rt-api',
                  'status': 'online', 'meta': {}, 'default_max_process': 5,
                  'default_max_timeout': 30, 'default_max_audio_duration': 60,
                  'endpoints': [], 'output_types': []}]},
        '/apis/%s')

    # ---------- 8. 清理 ----------
    removed, left = cleanup()
    print('cleanup removed', removed, 'left', left)
    check('收尾零残留', not left, left)

    print('\n==== API 契约验收：%d 项失败 / 用时 %.1fs ====' % (len(FAILS), time.time() - t0))
    for f in FAILS:
        print('FAILED:', f)
    return 1 if FAILS else 0


if __name__ == '__main__':
    sys.exit(main())
