# -*- coding: utf-8 -*-
"""INT-2226 M14：三端用例页关键路径实测（搜索 / 分组 / 导出 / 导入）。

环境：docker 真实库栈 http://localhost:15000；旧 Vue 8080 / React 15173 / 新 Vue 16173。
账号 dev_user / dev_password（主库唯一可用账号）。
写操作只落本 run 的「接线测试_INT2226_M14_」前缀数据，脚本结束前按 id 硬删。

跑法：
    <conda python> tests/ui/probe_int2226_m14_run.py --front react|uxvue|vue
产出：tests/ui/artifacts/int2226/m14_<front>.json + 浅/深色截图 + 下载文件。
"""
import argparse
import json
import os
import re
import sys
import time

import openpyxl
from playwright.sync_api import sync_playwright

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from tests.ui.utils.api_client import ApiClient  # noqa: E402

ART = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'artifacts', 'int2226')
os.makedirs(ART, exist_ok=True)
PREFIX = '接线测试_INT2226_M14'
GROUP = '话轮接管'          # 真实分组（只读筛选用）
AUDIO_ID = 2100            # 真实音频（H050…枪声_女）
AUDIO_NAME = 'HJ050_单一环境声音识别-无语义_枪声_女_silence_用户.wav'
BAD_AUDIO_ID = 999999

FRONTS = {
    'react': {'base': 'http://localhost:15173', 'hash': False, 'route': '/testcases'},
    'uxvue': {'base': 'http://localhost:16173', 'hash': True, 'route': '/#/testcases'},
    'vue': {'base': 'http://localhost:8080', 'hash': True, 'route': '/#/testcases'},
}


def make_fixture(path):
    """按 React「模板下载」得到的模板结构造带 2 合法 + 1 非法行的导入包。"""
    wb = openpyxl.load_workbook(os.path.join(ART, 'template.xlsx'))
    ws = wb['TestCases']
    ws.delete_rows(2, ws.max_row)
    rows = [
        (None, PREFIX + '_甲', 'M14 导入合法行 1', GROUP, None, None, 'api',
         None, None, None, None, None, None, None),
        (None, PREFIX + '_乙', 'M14 导入合法行 2', GROUP, None, None, 'api',
         None, None, None, None, None, None, None),
        (None, PREFIX + '_丙', 'M14 导入非法行（音频不存在）', GROUP, None, None, 'api',
         None, None, None, None, None, None, None),
    ]
    for r in rows:
        ws.append(r)
    wa = wb['AudioConfigs']
    wa.delete_rows(2, wa.max_row)
    for name, aid, spl in [(PREFIX + '_甲', AUDIO_ID, 65),
                           (PREFIX + '_乙', AUDIO_ID, 70),
                           (PREFIX + '_丙', BAD_AUDIO_ID, 65)]:
        wa.append((None, name, aid, AUDIO_NAME if aid == AUDIO_ID else None, spl, None, None, 1))
    for sheet in ('Dimensions', 'Tags', 'Groups', 'CaseTags'):
        ws2 = wb[sheet]
        if ws2.max_row > 1:
            ws2.delete_rows(2, ws2.max_row)
    wb.save(path)
    return path


def shot(page, tag):
    p = os.path.join(ART, 'm14_%s.png' % tag)
    page.screenshot(path=p, full_page=False)
    return p


def login(page, base, hash_route):
    page.goto(base + ('/#/login' if hash_route else '/login'), wait_until='domcontentloaded')
    time.sleep(3)
    page.locator('input[type=password]').first.fill('dev_password')
    page.locator('input[type=text], input:not([type])').first.fill('dev_user')
    btn = page.get_by_role('button', name=re.compile('登\\s*录'))
    (btn.first if btn.count() else page.locator('button[type=submit]').first).click()
    page.wait_for_function("() => !/login/i.test(location.href)", timeout=45000)
    time.sleep(3)


def set_theme(page, mode, hash_route):
    # React: iatp:theme:v1:mode；Vue: iatp:theme:v1
    key = 'iatp:theme:v1:mode' if not hash_route else 'iatp:theme:v1'
    page.evaluate("([k, m]) => localStorage.setItem(k, m)", [key, mode])
    page.reload(wait_until='domcontentloaded')
    time.sleep(4)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--front', required=True, choices=list(FRONTS))
    ap.add_argument('--skip-import', action='store_true')
    args = ap.parse_args()
    cfg = FRONTS[args.front]
    base, hash_route, route = cfg['base'], cfg['hash'], cfg['route']
    out = {'front': args.front, 'base': base, 'steps': {}, 'api': [], 'shots': []}

    api = ApiClient('http://localhost:15000')
    api.login('dev_user', 'dev_password')
    health = api.health()
    out['health'] = health
    assert health.get('database') == 'intelligent_audio_test', health

    with sync_playwright() as p:
        b = p.chromium.launch(headless=True)
        ctx = b.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True)
        page = ctx.new_page()
        resp = []

        def on_resp(r):
            if '/api/v1/' in r.url:
                rec = {'url': r.url.replace(base, ''), 'status': r.status, 'method': r.request.method}
                if any(k in r.url for k in ('/testcases/import', '/testcases/export')):
                    try:
                        rec['body'] = r.text()[:1500]
                    except Exception:
                        pass
                resp.append(rec)

        page.on('response', on_resp)
        login(page, base, hash_route)
        page.goto(base + route, wait_until='domcontentloaded')
        time.sleep(6)

        # ── 环境：用例总数以界面与接口对照 ────────────────────
        total_api = api.request('GET', '/testcases?page=1&per_page=1').get('data', {}).get('total')
        body = page.inner_text('body')
        out['steps']['landing'] = {
            'url': page.url, 'api_total': total_api,
            'heading_ok': '用例' in body,
            'shot_light': shot(page, '%s_light' % args.front),
        }

        # ── 搜索 ──────────────────────────────────────────
        search = page.locator('input[placeholder*="搜索"]').first
        search.click()
        search.fill('接线测试_INT2032')
        page.keyboard.press('Enter')
        time.sleep(4)
        txt_after = page.inner_text('body')
        out['steps']['search'] = {
            'keyword': '接线测试_INT2032',
            'visible_matches': page.locator('text=接线测试_INT2032').count(),
            'body_has_keyword': '接线测试_INT2032' in txt_after,
            'api_calls': [r for r in resp if '/testcases' in r['url'] and 'keyword' in r['url']][-3:],
            'shot': shot(page, '%s_search' % args.front),
        }
        search.fill('zzz_no_hit_2226')
        page.keyboard.press('Enter')
        time.sleep(4)
        out['steps']['search']['empty_body_has_keyword'] = 'zzz_no_hit_2226' in page.inner_text('body')
        out['steps']['search']['empty_hit_rows'] = page.locator('text=zzz_no_hit_2226').count()
        out['steps']['search']['empty_state_text'] = page.inner_text('body')[-300:]
        search.fill('')
        page.keyboard.press('Enter')
        time.sleep(4)

        # ── 分组筛选 ───────────────────────────────────────
        group_step = {'group': GROUP}
        api_group_total = None
        groups = api.request('GET', '/groups?page=1&per_page=100').get('data', {}).get('items') or []
        for g in groups:
            if g.get('name') == GROUP:
                api_group_total = g.get('test_case_count')
        group_step['api_group_count'] = api_group_total
        if args.front == 'react':
            page.get_by_role('button', name=re.compile('^分组')).first.click()
            time.sleep(1.5)
            opt = page.locator('[role=menu] >> text=%s' % GROUP)
            if not opt.count():
                opt = page.get_by_text(GROUP, exact=True)
            group_step['option_count'] = opt.count()
            if opt.count():
                opt.last.click()
            time.sleep(4)
            group_step['api_filtered'] = [r for r in resp if 'group_id=' in r['url']][-2:]
        else:
            # Vue 两端的「分组」是徽标直选（不是下拉）
            gb = page.locator('button.badge-filter__badge', has_text=GROUP)
            group_step['option_count'] = gb.count()
            if gb.count():
                gb.first.click()
            time.sleep(4)
            group_step['api_filtered'] = [r for r in resp
                                          if ('group_id=' in r['url'] or 'group=' in r['url'])][-2:]
        body_g = page.inner_text('body')
        m = (re.search(re.escape(GROUP) + r'\s*(\d+)\s*条', body_g) or
             re.search(re.escape(GROUP) + r'\s*(\d+)', body_g))
        group_step['ui_group_count'] = int(m.group(1)) if m else None
        group_step['count_ok'] = (api_group_total is not None and
                                  group_step['ui_group_count'] == api_group_total)
        group_step['shot'] = shot(page, '%s_group' % args.front)
        out['steps']['group'] = group_step

        # ── 导出（按当前分组筛选）────────────────────────────
        exp = {'format': None}
        try:
            if args.front == 'react':
                page.get_by_role('button', name=re.compile('^导出')).first.click()
                time.sleep(2)
                dlg = page.locator('[role=dialog]').filter(has_text='导出用例').first
                exp['dialog_text'] = dlg.inner_text()[:600]
                exp['preview'] = re.search(r'将导出\s*(\d+)\s*条', exp['dialog_text']).group(1) \
                    if re.search(r'将导出\s*(\d+)\s*条', exp['dialog_text']) else None
                exp['format'] = 'zip'
                with page.expect_download(timeout=180000) as dl:
                    dlg.get_by_role('button', name='开始导出').click()
                f = os.path.join(ART, 'm14_%s_export.zip' % args.front)
                dl.value.save_as(f)
                exp['file'] = f
                exp['size'] = os.path.getsize(f)
            else:
                page.get_by_role('button', name='导出用例').first.click()
                time.sleep(2)
                dlg = page.locator('.modal-overlay.active').first
                exp['dialog_text'] = dlg.inner_text()[:900]
                # 勾选测试组（弹窗里 id=group-<名称>）
                cb = dlg.locator('#group-%s' % GROUP)
                if not cb.count():
                    cb = dlg.locator('label', has_text=GROUP).first
                if cb.count():
                    cb.first.check()
                time.sleep(2)
                exp['preview'] = re.search(r'预计导出[：:]\s*(\d+)\s*条', dlg.inner_text())
                exp['preview'] = exp['preview'].group(1) if exp['preview'] else None
                exp['format'] = 'xlsx'
                with page.expect_download(timeout=180000) as dl:
                    dlg.get_by_role('button', name=re.compile('^导出$')).first.click()
                f = os.path.join(ART, 'm14_%s_export.xlsx' % args.front)
                dl.value.save_as(f)
                exp['file'] = f
                exp['size'] = os.path.getsize(f)
            exp['ok'] = True
        except Exception as exc:
            exp['ok'] = False
            exp['error'] = '%s: %s' % (type(exc).__name__, str(exc)[:300])
            try:
                exp['shot'] = shot(page, '%s_export_fail' % args.front)
            except Exception:
                pass
        exp['api'] = [r for r in resp if '/testcases/export' in r['url']][-3:]
        out['steps']['export'] = exp

        # ── 导入 ─────────────────────────────────────────
        if not args.skip_import:
            imp = {'mode': 'file' if args.front == 'react' else 'audio-import-route'}
            try:
                page.goto(base + route, wait_until='domcontentloaded')
                time.sleep(4)
                page.get_by_role('button', name='导入用例').first.click()
                time.sleep(3)
                if args.front == 'react':
                    dlg = page.locator('[role=dialog]').filter(has_text='导入用例').first
                    imp['dialog_text'] = dlg.inner_text()[:500]
                    fixture = make_fixture(os.path.join(ART, 'm14_import_fixture.xlsx'))
                    dlg.locator('input[type=file]').set_input_files(fixture)
                    time.sleep(8)
                    imp['preview_text'] = dlg.inner_text()[:700] if dlg.count() else ''
                    goto_btn = dlg.get_by_role('button', name=re.compile('仍导入|开始导入|确认导入'))
                    imp['import_button'] = goto_btn.first.inner_text() if goto_btn.count() else None
                    if goto_btn.count():
                        goto_btn.first.click()
                    time.sleep(30)
                    imp['after_text'] = page.locator('body').inner_text()[-1200:]
                    imp['dialog_after'] = dlg.inner_text()[:1200] if dlg.count() else ''
                else:
                    imp['url_after_click'] = page.url
                    modal = page.locator('.modal-overlay.active').first
                    imp['modal_text'] = modal.inner_text()[:700] if modal.count() else ''
                    imp['create_case_checked'] = False
                    try:
                        imp['checkboxes'] = modal.evaluate("""el => [...el.querySelectorAll('input[type=checkbox]')]
                          .map(c => ({checked: c.checked, id: c.id,
                                      label: (c.closest('label') || {}).innerText || ''}))""")
                        imp['create_case_checked'] = any(
                            '生成测试用例' in (c.get('label') or '') and c.get('checked')
                            for c in imp['checkboxes'])
                    except Exception:
                        pass
                    imp['opens_file_import'] = 'xlsx' in imp['modal_text'] or '.zip' in imp['modal_text']
                imp['shot'] = shot(page, '%s_import' % args.front)
                imp['api'] = [r for r in resp if '/testcases/import' in r['url']][-5:]
                imp['ok'] = True
            except Exception as exc:
                imp['ok'] = False
                imp['error'] = '%s: %s' % (type(exc).__name__, str(exc)[:300])
                try:
                    imp['shot'] = shot(page, '%s_import_fail' % args.front)
                except Exception:
                    pass
            out['steps']['import'] = imp

        # ── 深色截图 ──────────────────────────────────────
        try:
            page.goto(base + route, wait_until='domcontentloaded')
            time.sleep(4)
            set_theme(page, 'dark', hash_route)
            page.goto(base + route, wait_until='domcontentloaded')
            time.sleep(5)
            out['steps']['dark'] = {
                'html_class': page.evaluate("() => document.documentElement.className"),
                'shot': shot(page, '%s_dark' % args.front),
            }
        except Exception as exc:
            out['steps']['dark'] = {'error': str(exc)[:200]}
        set_theme(page, 'light', hash_route)

        out['api'] = resp[-60:]
        b.close()

    # ── 清理本 run 导入的数据 ─────────────────────────────
    created = api.request('GET', '/testcases?keyword=' + PREFIX + '&page=1&per_page=50')
    items = (created.get('data') or {}).get('items') or []
    out['cleanup'] = {'found': [(i.get('id'), i.get('name')) for i in items]}
    for i in items:
        try:
            api.delete_testcase(i.get('id'), purge=True)
        except Exception as exc:
            out['cleanup'].setdefault('errors', []).append('%s: %s' % (i.get('id'), exc))
    left = api.request('GET', '/testcases?keyword=' + PREFIX + '&page=1&per_page=50')
    out['cleanup']['remaining'] = (left.get('data') or {}).get('total')

    path = os.path.join(ART, 'm14_%s.json' % args.front)
    with open(path, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    print(json.dumps(out, ensure_ascii=False, indent=1)[:6000], flush=True)
    print('WROTE', path, flush=True)


if __name__ == '__main__':
    main()
